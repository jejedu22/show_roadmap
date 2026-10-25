from __future__ import annotations

import io
from collections import defaultdict
from urllib.parse import quote as urlquote

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import RedirectResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy.orm import Session

from ..auth import accessible_company, csrf_protect, current_user
from ..db import get_db
from ..models import (Day, Event, ExtraCost, Leg, Lodging, Participation, Person, RoadmapSection, ScheduleItem,
                      TruckLeg, User)
from ..services import (compute_quote, event_overlaps, fill_default_days, person_conflicts, person_costs,
                        truck_scale)
from ..web import p_date, p_float, p_int, p_str, render
from .core import user_companies

router = APIRouter(prefix="/events", dependencies=[Depends(csrf_protect)])

TABS = [("", "Aperçu"), ("team", "Équipe"), ("stay", "Séjour"), ("travel", "Voyages"), ("quote", "Devis"),
        ("roadmap", "Feuille de route")]


def get_event(db: Session, user: User, eid: int) -> Event:
    ev = db.get(Event, eid)
    if not ev:
        raise HTTPException(404, "Date introuvable")
    accessible_company(db, user, ev.show.company_id)
    return ev


def page(request, name, user, ev, db, tab, **ctx):
    return render(request, name, user=user, ev=ev, company=ev.show.company, tab=tab, tabs=TABS,
                  nav_companies=user_companies(db, user), **ctx)


def back(eid: int, tab: str = "") -> RedirectResponse:
    return RedirectResponse(f"/events/{eid}" + (f"/{tab}" if tab else ""), status_code=303)


# ---------- Aperçu / édition ----------

@router.get("/{eid}")
def overview(eid: int, request: Request, user: User = Depends(current_user), db: Session = Depends(get_db)):
    ev = get_event(db, user, eid)
    return page(request, "event_overview.html", user, ev, db, "", overlaps=event_overlaps(db, ev),
                conflicts=person_conflicts(db, ev), quote=compute_quote(ev))


@router.post("/{eid}")
async def update_event(eid: int, request: Request, user: User = Depends(current_user), db: Session = Depends(get_db)):
    ev = get_event(db, user, eid)
    f = await request.form()
    ev.venue = p_str(f.get("venue")) or ev.venue
    ev.city, ev.address = p_str(f.get("city")), p_str(f.get("address"))
    ev.venue_contact, ev.notes = f.get("venue_contact") or "", f.get("notes") or ""
    ev.status = f.get("status") if f.get("status") in ("option", "confirmed", "cancelled") else ev.status
    for k in ("first_date", "last_date"):
        setattr(ev, k, p_date(f.get(k)) or getattr(ev, k))
    for k in ("tech_arrival", "artistic_arrival", "departure"):
        setattr(ev, k, p_date(f.get(k)))
    ev.performances = max(1, p_int(f.get("performances"), 1))
    if ev.last_date < ev.first_date:
        raise HTTPException(400, "La dernière date précède la première")
    db.commit()
    return back(eid)


@router.post("/{eid}/delete")
def delete_event(eid: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    ev = get_event(db, user, eid)
    cid = ev.show.company_id
    db.delete(ev)
    db.commit()
    return RedirectResponse(f"/companies/{cid}", status_code=303)


@router.post("/{eid}/duplicate")
def duplicate_event(eid: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    """Copie la date (sans voyages) pour préparer une nouvelle date du même spectacle."""
    ev = get_event(db, user, eid)
    new = Event(**{c.name: getattr(ev, c.name) for c in Event.__table__.columns if c.name != "id"})
    new.venue = f"{ev.venue} (copie)"
    db.add(new)
    db.flush()
    for lg in ev.lodgings:
        db.add(Lodging(event_id=new.id, name=lg.name, address=lg.address, phone=lg.phone, details=lg.details))
    for p in ev.participations:
        np_ = Participation(event_id=new.id, person_id=p.person_id, role=p.role, room_type=p.room_type,
                            departure_city=p.departure_city)
        db.add(np_)
        db.flush()
        fill_default_days(db, np_)
    db.commit()
    return back(new.id)


# ---------- Équipe ----------

@router.get("/{eid}/team")
def team(eid: int, request: Request, user: User = Depends(current_user), db: Session = Depends(get_db)):
    ev = get_event(db, user, eid)
    used = {p.person_id for p in ev.participations}
    available = [p for p in ev.show.company.people if p.id not in used]
    return page(request, "event_team.html", user, ev, db, "team", available=available,
                conflicts=person_conflicts(db, ev))


@router.post("/{eid}/team/add")
async def team_add(eid: int, request: Request, user: User = Depends(current_user), db: Session = Depends(get_db)):
    ev = get_event(db, user, eid)
    f = await request.form()
    ids = [int(x) for x in f.getlist("person_id") if str(x).isdigit()]
    existing = {p.person_id for p in ev.participations}
    for pid in ids:
        person = db.get(Person, pid)
        if not person or person.company_id != ev.show.company_id or pid in existing:
            continue
        part = Participation(event_id=ev.id, person_id=pid, role=person.function, departure_city=person.home_city)
        db.add(part)
        db.flush()
        db.refresh(part)
        fill_default_days(db, part)
    db.commit()
    return back(eid, "team")


@router.post("/{eid}/team/save")
async def team_save(eid: int, request: Request, user: User = Depends(current_user), db: Session = Depends(get_db)):
    ev = get_event(db, user, eid)
    f = await request.form()
    for p in ev.participations:
        p.role = p_str(f.get(f"role_{p.id}"))
        p.room_type = p_str(f.get(f"room_{p.id}")) or "single"
        p.departure_city = p_str(f.get(f"city_{p.id}"))
        lid = p_int(f.get(f"lodging_{p.id}"))
        p.lodging_id = lid if any(l.id == lid for l in ev.lodgings) else None
        p.note = p_str(f.get(f"note_{p.id}"))
    db.commit()
    return back(eid, "team")


@router.post("/{eid}/team/{part_id}/remove")
def team_remove(eid: int, part_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    ev = get_event(db, user, eid)
    part = db.get(Participation, part_id)
    if not part or part.event_id != ev.id:
        raise HTTPException(404)
    db.delete(part)
    db.commit()
    return back(eid, "team")


# ---------- Séjour : nuitées / repas / petits-déjeuners ----------

@router.get("/{eid}/stay")
def stay(eid: int, request: Request, user: User = Depends(current_user), db: Session = Depends(get_db)):
    ev = get_event(db, user, eid)
    return page(request, "event_stay.html", user, ev, db, "stay", dates=ev.span_dates)


@router.post("/{eid}/stay")
async def stay_save(eid: int, request: Request, user: User = Depends(current_user), db: Session = Depends(get_db)):
    ev = get_event(db, user, eid)
    f = await request.form()
    for part in ev.participations:
        existing = {d.day: d for d in part.days}
        for d in ev.span_dates:
            key = d.isoformat()
            vals = (min(1, max(0, p_int(f.get(f"n_{part.id}_{key}")))), min(2, max(0, p_int(f.get(f"m_{part.id}_{key}")))),
                    min(1, max(0, p_int(f.get(f"b_{part.id}_{key}")))))
            row = existing.get(d)
            if row:
                row.night, row.meals, row.breakfast = vals
            elif any(vals):
                db.add(Day(participation_id=part.id, day=d, night=vals[0], meals=vals[1], breakfast=vals[2]))
    db.commit()
    return back(eid, "stay")


@router.post("/{eid}/stay/reset")
def stay_reset(eid: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    ev = get_event(db, user, eid)
    for part in ev.participations:
        for d in list(part.days):
            db.delete(d)
        db.flush()
        db.refresh(part)
        fill_default_days(db, part)
    db.commit()
    return back(eid, "stay")


# ---------- Voyages, devis, feuille de route : CRUD générique des lignes ----------
# kind -> (modèle, champs {nom: parseur}, onglet de retour)
SPEC = {
    "lodging": (Lodging, {"name": p_str, "address": p_str, "phone": p_str, "details": lambda v: v or ""}, "stay"),
    "leg": (Leg, {"participation_id": p_int, "direction": p_str, "day": p_date, "origin": p_str, "destination": p_str,
                  "dep_time": p_str, "arr_time": p_str, "mode": p_str, "reference": p_str, "price": p_float,
                  "booked_on": p_date, "note": p_str}, "travel"),
    "truck": (TruckLeg, {"label": p_str, "km": p_float}, "quote"),
    "extra": (ExtraCost, {"label": p_str, "amount": p_float, "kind": p_str}, "quote"),
    "section": (RoadmapSection, {"title": p_str, "body": lambda v: v or "", "position": p_int}, "roadmap"),
    "schedule": (ScheduleItem, {"day": p_date, "time": p_str, "label": p_str}, "roadmap"),
}


def _owned(db: Session, ev: Event, kind: str, item_id: int):
    model = SPEC[kind][0]
    obj = db.get(model, item_id)
    if not obj:
        raise HTTPException(404)
    owner = obj.participation.event_id if kind == "leg" else obj.event_id
    if owner != ev.id:
        raise HTTPException(404)
    return obj


@router.post("/{eid}/item/{kind}")
async def item_save(eid: int, kind: str, request: Request, user: User = Depends(current_user),
                    db: Session = Depends(get_db)):
    if kind not in SPEC:
        raise HTTPException(404)
    ev = get_event(db, user, eid)
    model, fields, tab = SPEC[kind]
    f = await request.form()
    vals = {k: parse(f.get(k)) for k, parse in fields.items() if k in f}
    if kind == "leg":
        part = db.get(Participation, vals.get("participation_id", 0))
        if not part or part.event_id != ev.id:
            raise HTTPException(400, "Participant invalide")
        if vals.get("direction") not in ("outbound", "return"):
            vals["direction"] = "outbound"
        if not vals.get("day"):
            raise HTTPException(400, "Date du trajet obligatoire")
    if kind == "extra" and vals.get("kind") not in ("team_transport", "other"):
        vals["kind"] = "other"
    if kind in ("schedule",) and not vals.get("day"):
        raise HTTPException(400, "Date obligatoire")
    item_id = p_int(f.get("id"))
    if item_id:
        obj = _owned(db, ev, kind, item_id)
        if kind == "leg":
            vals.pop("participation_id", None)
        for k, v in vals.items():
            setattr(obj, k, v)
    else:
        obj = model(**vals) if kind == "leg" else model(event_id=ev.id, **vals)
        db.add(obj)
    db.commit()
    return back(eid, tab)


@router.post("/{eid}/item/{kind}/{item_id}/delete")
def item_delete(eid: int, kind: str, item_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    if kind not in SPEC:
        raise HTTPException(404)
    ev = get_event(db, user, eid)
    db.delete(_owned(db, ev, kind, item_id))
    db.commit()
    return back(eid, SPEC[kind][2])


@router.get("/{eid}/travel")
def travel(eid: int, request: Request, user: User = Depends(current_user), db: Session = Depends(get_db)):
    ev = get_event(db, user, eid)
    return page(request, "event_travel.html", user, ev, db, "travel", costs=person_costs(ev))


@router.get("/{eid}/quote")
def quote(eid: int, request: Request, user: User = Depends(current_user), db: Session = Depends(get_db)):
    ev = get_event(db, user, eid)
    km = sum(t.km for t in ev.trucks)
    return page(request, "event_quote.html", user, ev, db, "quote", q=compute_quote(ev),
                scale=truck_scale(km, ev.performances))


@router.post("/{eid}/quote")
async def quote_save(eid: int, request: Request, user: User = Depends(current_user), db: Session = Depends(get_db)):
    ev = get_event(db, user, eid)
    f = await request.form()
    for k in ("meal_rate", "breakfast_rate", "night_rate", "truck_rate", "vat_rate", "fee", "lighting", "posters"):
        setattr(ev, k, p_float(f.get(k)))
    db.commit()
    return back(eid, "quote")


@router.get("/{eid}/roadmap")
def roadmap_edit(eid: int, request: Request, user: User = Depends(current_user), db: Session = Depends(get_db)):
    ev = get_event(db, user, eid)
    return page(request, "event_roadmap_edit.html", user, ev, db, "roadmap")


# ---------- Documents imprimables ----------

def _travel_by_day(ev: Event):
    days = defaultdict(list)
    for part in ev.participations:
        for lg in part.legs:
            days[lg.day].append((part, lg))
    return {d: sorted(v, key=lambda x: (x[1].dep_time or "99", x[0].person.last_name)) for d, v in sorted(days.items())}


@router.get("/{eid}/docs/{doc}")
def document(eid: int, doc: str, request: Request, user: User = Depends(current_user), db: Session = Depends(get_db)):
    ev = get_event(db, user, eid)
    templates = {"roadmap": "doc_roadmap.html", "rooming": "doc_rooming.html", "perdiem": "doc_perdiem.html",
                 "travelcost": "doc_travelcost.html", "contacts": "doc_contacts.html", "quote": "doc_quote.html"}
    if doc not in templates:
        raise HTTPException(404)
    return render(request, templates[doc], ev=ev, dates=ev.span_dates, travel_by_day=_travel_by_day(ev),
                  costs=person_costs(ev), q=compute_quote(ev), user=user)


@router.get("/{eid}/export.xlsx")
def export_xlsx(eid: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    ev = get_event(db, user, eid)
    wb = Workbook()
    bold = Font(bold=True)

    def sheet(title, header, rows, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.font = bold
        for r in rows:
            ws.append(r)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(40, max(10, *(len(str(c.value or "")) for c in col)))
        return ws

    q = compute_quote(ev)
    sheet("Devis", ["Poste", "Montant HT (€)"],
          [["Cachet", q.fee]] + [[a, b] for a, b in q.lines] +
          [["Total frais annexes", q.annex_total], ["Total HT", q.total_ht],
           [f"TVA {ev.vat_rate:g} %", q.vat], ["Total TTC", q.total_ttc]], first=True)
    dates = ev.span_dates
    sheet("Rooming", ["Nom", "Fonction", "Chambre", "Hébergement"] + [d.strftime("%d/%m") for d in dates] + ["Total"],
          [[p.person.full_name, p.role, p.room_type, p.lodging.name if p.lodging else ""] +
           [(p.day(d).night if p.day(d) else 0) for d in dates] + [sum(x.night for x in p.days)]
           for p in ev.participations])
    sheet("Défraiements", ["Nom", "Repas", "Petits-déj.", "Total (€)"],
          [[c.part.person.full_name, c.meals, c.breakfasts, c.per_diem] for c in person_costs(ev)])
    sheet("Voyages", ["Nom", "Sens", "Date", "De", "À", "Départ", "Arrivée", "Mode", "Référence", "Prix (€)"],
          [[p.person.full_name, "Aller" if l.direction == "outbound" else "Retour", l.day, l.origin, l.destination,
            l.dep_time, l.arr_time, l.mode, l.reference, l.price] for p in ev.participations for l in p.legs])
    sheet("Coordonnées", ["Nom", "Fonction", "Téléphone", "E-mail", "Régime"],
          [[p.person.full_name, p.role, p.person.phone, p.person.email, p.person.diet] for p in ev.participations])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"{ev.show.title}_{ev.venue}".replace(" ", "_").replace("/", "-")
    ascii_name = name.encode("ascii", "ignore").decode() or "export"
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=\"{ascii_name}.xlsx\"; filename*=UTF-8''{urlquote(name)}.xlsx"})
